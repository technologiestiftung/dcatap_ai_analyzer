import pandas as pd
from pathlib import Path


def compare_detailed_scores(old_score_table_path, new_score_table_path, 
                           old_metadata_path, new_metadata_path):
    """Create detailed per-score comparison for each organization.
    
    Args:
        old_score_table_path: Path to old Score_Tabelle.xlsx
        new_score_table_path: Path to new Score_Tabelle.xlsx
        old_metadata_path: Path to old metadata_analysis file
        new_metadata_path: Path to new metadata_analysis file
    
    Returns:
        DataFrame with detailed per-score comparison
    """
    # Read metadata files to get individual scores
    old_meta = pd.read_excel(old_metadata_path)
    new_meta = pd.read_excel(new_metadata_path)
    
    # Load harmonization to match organization names
    harmonization_path = Path("_results/Harmonize_Authors_Arrays.xlsx")
    if harmonization_path.exists():
        from create_org_scores import load_author_harmonization, harmonize_authors
        
        harmonization_df = pd.read_excel(harmonization_path)
        author_mapping = load_author_harmonization(harmonization_path)
        
        old_meta = harmonize_authors(old_meta, author_mapping, harmonization_df)
        new_meta = harmonize_authors(new_meta, author_mapping, harmonization_df)
    
    # Get list of all organizations from Score_Tabelle files
    old_scores = pd.read_excel(old_score_table_path)
    new_scores = pd.read_excel(new_score_table_path)
    
    all_orgs = set(old_scores['Zuständige Stelle'].tolist()) | \
               set(new_scores['Zuständige Stelle'].tolist())
    
    # Score columns to compare (all 6 scores)
    score_cols = [
        'dateninhalt_score',
        'methodik_score',
        'datenqualitaet_score',
        'geographie_score',
        'tag_qualitaet_score',
        'referenz_score'
    ]
    
    # Build detailed comparison
    results = []
    
    for org in sorted(all_orgs):
        row = {'Organization': org}
        
        # Get old data
        old_org_data = old_meta[old_meta['author_harmonized'] == org] if 'author_harmonized' in old_meta.columns else old_meta[old_meta['author'] == org]
        # Get new data
        new_org_data = new_meta[new_meta['author_harmonized'] == org] if 'author_harmonized' in new_meta.columns else new_meta[new_meta['author'] == org]
        
        row['Old Datasets'] = len(old_org_data)
        row['New Datasets'] = len(new_org_data)
        row['Dataset Change'] = row['New Datasets'] - row['Old Datasets']
        
        # Calculate averages for each score
        for score_col in score_cols:
            score_name = score_col.replace('_score', '')
            
            if len(old_org_data) > 0 and score_col in old_org_data.columns:
                old_avg = old_org_data[score_col].mean()
            else:
                old_avg = None
            
            if len(new_org_data) > 0 and score_col in new_org_data.columns:
                new_avg = new_org_data[score_col].mean()
            else:
                new_avg = None
            
            row[f'{score_name}_old'] = old_avg
            row[f'{score_name}_new'] = new_avg
            
            if old_avg is not None and new_avg is not None:
                row[f'{score_name}_change'] = new_avg - old_avg
            else:
                row[f'{score_name}_change'] = None
        
        # Get total scores from Score_Tabelle
        old_total = old_scores[old_scores['Zuständige Stelle'] == org]['total_score'].iloc[0] if org in old_scores['Zuständige Stelle'].values else None
        new_total = new_scores[new_scores['Zuständige Stelle'] == org]['total_score'].iloc[0] if org in new_scores['Zuständige Stelle'].values else None
        
        row['total_score_old'] = old_total
        row['total_score_new'] = new_total
        
        if old_total is not None and new_total is not None:
            row['total_score_change'] = new_total - old_total
        else:
            row['total_score_change'] = None
        
        results.append(row)
    
    return pd.DataFrame(results)


def create_formatted_excel(df, output_path):
    """Create a nicely formatted Excel file with the detailed comparison."""
    
    # Reorder columns for better readability
    column_order = ['Organization', 'Old Datasets', 'New Datasets', 'Dataset Change']
    
    # All 6 scores
    all_score_names = ['dateninhalt', 'methodik', 'datenqualitaet', 'geographie', 
                       'tag_qualitaet', 'referenz']
    
    for score in all_score_names:
        column_order.extend([f'{score}_old', f'{score}_new', f'{score}_change'])
    
    column_order.extend(['total_score_old', 'total_score_new', 'total_score_change'])
    
    # Reorder
    df = df[column_order]
    
    # Sort by total score change (biggest improvements first)
    df = df.sort_values('total_score_change', ascending=False, na_position='last')
    
    # Write to Excel
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Detailed Score Comparison', index=False)
        
        # Get workbook and worksheet
        workbook = writer.book
        worksheet = writer.sheets['Detailed Score Comparison']
        
        # Format header row
        from openpyxl.styles import Font, PatternFill, Alignment
        
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Color code changes (green for positive, red for negative)
        green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        
        change_cols = [col for col in df.columns if col.endswith('_change')]
        change_col_indices = [df.columns.get_loc(col) + 1 for col in change_cols]  # +1 for Excel 1-indexing
        
        for row_idx in range(2, len(df) + 2):  # Start from row 2 (after header)
            for col_idx in change_col_indices:
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    if cell.value > 0:
                        cell.fill = green_fill
                    elif cell.value < 0:
                        cell.fill = red_fill
                    
                    # Format to 3 decimal places
                    cell.number_format = '0.000'
        
        # Format score columns to 3 decimal places
        score_col_names = [col for col in df.columns if any(s in col for s in ['_old', '_new']) and not col.startswith('Old') and not col.startswith('New')]
        score_col_indices = [df.columns.get_loc(col) + 1 for col in score_col_names]
        
        for row_idx in range(2, len(df) + 2):
            for col_idx in score_col_indices:
                cell = worksheet.cell(row=row_idx, column=col_idx)
                if cell.value is not None and isinstance(cell.value, (int, float)):
                    cell.number_format = '0.000'
    
    print(f"✓ Saved detailed comparison to: {output_path}")


def print_summary(df):
    """Print a summary of the detailed comparison."""
    print("\n" + "="*80)
    print("DETAILED PER-SCORE COMPARISON SUMMARY")
    print("="*80)
    
    # Organizations with data in both periods
    existing = df[df['total_score_old'].notna() & df['total_score_new'].notna()]
    
    if len(existing) == 0:
        print("\nNo organizations with data in both periods.")
        return
    
    # All 6 scores
    score_names = ['dateninhalt', 'methodik', 'datenqualitaet', 'geographie',
                   'tag_qualitaet', 'referenz']
    
    print(f"\nOrganizations analyzed: {len(existing)}")
    print("\nAverage changes per score type:")
    print("-" * 80)
    
    for score in score_names:
        change_col = f'{score}_change'
        if change_col in df.columns:
            avg_change = existing[change_col].mean()
            improved = (existing[change_col] > 0).sum()
            declined = (existing[change_col] < 0).sum()
            unchanged = (existing[change_col] == 0).sum()
            
            score_display = score.replace('data', 'daten').replace('quality', 'qualität')
            print(f"  {score_display:20s}: {avg_change:+.3f} avg "
                  f"(↑{improved} ↓{declined} ={unchanged})")
    
    print(f"\n  {'total_score':20s}: {existing['total_score_change'].mean():+.3f} avg")
    
    # Top improvers per score
    print("\n" + "-"*80)
    print("TOP 5 IMPROVERS PER SCORE TYPE:")
    print("-"*80)
    
    for score in score_names:
        change_col = f'{score}_change'
        if change_col in df.columns:
            top_5 = existing.nlargest(5, change_col)
            score_display = score.replace('data', 'daten').replace('quality', 'qualität')
            print(f"\n{score_display.upper()}:")
            for _, row in top_5.iterrows():
                if pd.notna(row[change_col]) and row[change_col] > 0:
                    print(f"  {row['Organization']:45s} {row[change_col]:+.3f}")


if __name__ == "__main__":
    # Paths
    old_score_table = Path("_results/Score_Tabelle.xlsx")
    new_score_table = Path("_results/open_data_ag_20251204/Score_Tabelle20251204.xlsx")
    
    # Use the most recent February metadata (should match the old Score_Tabelle)
    old_metadata = Path("_results/metadata_analysis_20250224.xlsx")
    new_metadata = Path("_results/metadata_analysis_20251204.xlsx")
    
    output_path = Path("_results/Score_Comparison_Detailed.xlsx")
    
    print("Creating detailed per-score comparison...")
    print(f"  Old data: {old_metadata.name}")
    print(f"  New data: {new_metadata.name}")
    
    # Create detailed comparison
    detailed_df = compare_detailed_scores(
        old_score_table, 
        new_score_table,
        old_metadata,
        new_metadata
    )
    
    # Save to Excel with formatting
    create_formatted_excel(detailed_df, output_path)
    
    # Print summary
    print_summary(detailed_df)
    
    print("\n" + "="*80)
    print("DETAILED COMPARISON COMPLETE")
    print("="*80)
    print(f"\nResults saved to: {output_path}")
    print("Open the Excel file to see color-coded changes:")
    print("  - Green: Improvements")
    print("  - Red: Declines")
